import time
from flask import Blueprint, request, jsonify, current_app
from werkzeug.utils import secure_filename
import os
import openpyxl
import traceback
from sqlalchemy import create_engine, inspect, MetaData, Table
from sqlalchemy.orm import sessionmaker
from sqlalchemy.ext.declarative import declarative_base
from sqlalchemy import Column, Integer, String, Boolean
from .Db import db
import re


Base = declarative_base()
upload_blueprint = Blueprint('upload', __name__)

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in current_app.config['ALLOWED_EXTENSIONS']

def drop_table_if_exists(table_name):
    """
    Slett tabellen hvis den allerede finnes.
    """
    engine = db.engine
    inspector = inspect(engine)
    
    if table_name in inspector.get_table_names():
        metadata = MetaData()
        table = Table(table_name, metadata, autoload_with=engine)
        table.drop(engine)
        print(f"Tabellen {table_name} ble slettet.")
        # Fjern fra metadata cache
        if table_name in Base.metadata.tables:
            Base.metadata.remove(Base.metadata.tables[table_name])
    else:
        print(f"Tabellen {table_name} finnes ikke.")


def create_dynamic_model(table_name, headers):
    """
    Create a SQLAlchemy model dynamically based on the headers from the Excel file.
    """
    # Sanitize headers to be valid Python identifiers
    sanitized_headers = [re.sub(r'\W|^(?=\d)', '_', header) for header in headers]

    # Create columns dynamically based on headers
    columns = {
        '__tablename__': table_name,
        'id': Column(Integer, primary_key=True, autoincrement=True),  # Auto-generated primary key
        'ischecked': Column(Boolean, default=False),  # Add ischecked column with default False
    }

    for header, sanitized_header in zip(headers, sanitized_headers):
        # Add columns dynamically based on the sanitized header names
        columns[sanitized_header] = Column(String(255), nullable=True)

    # Create and return a dynamic model class
    DynamicModel = type(table_name, (Base,), columns)
    
    return DynamicModel


@upload_blueprint.route('/upload-excel', methods=['POST'])
def upload_excel():
    if 'file' not in request.files:
        return jsonify({'error': 'No file part'}), 400

    file = request.files['file']
    if file.filename == '' or not allowed_file(file.filename):
        return jsonify({'error': 'No selected file or invalid file type'}), 400

    filename = secure_filename(file.filename)
    file_path = os.path.join(current_app.config['UPLOAD_FOLDER'], filename)
    os.makedirs(current_app.config['UPLOAD_FOLDER'], exist_ok=True)
    file.save(file_path)

    try:
        # Load the Excel file
        workbook = openpyxl.load_workbook(file_path)
        sheet = workbook.active

        # Extract headers
        headers = [str(cell.value).strip() if cell.value else f"Column_{i}" for i, cell in enumerate(sheet[1], start=1)]
        print(f"Original Headers: {headers}")  # Debugging: Print original headers
        
        # Exclude headers that start with "Column"
        headers = [header for header in headers if not header.startswith("Column")]
        print(f"Filtered Headers: {headers}")  # Debugging: Print filtered headers

        # Sanitize headers to be valid Python identifiers
        sanitized_headers = [re.sub(r'\W|^(?=\d)', '_', header) for header in headers]
        print(f"Sanitized Headers: {sanitized_headers}")  # Debugging: Print sanitized headers

        # Slett tabellen hvis den allerede finnes
        table_name = 'imported_table'
        drop_table_if_exists(table_name)
        
        # Create a dynamic model
        DynamicModel = create_dynamic_model(table_name, sanitized_headers)
        DynamicModel.metadata.create_all(db.engine)  # Opprett tabellen
        
        # Connect to the database and create the table
        # Explicitly create the table
        Base.metadata.create_all(db.engine)  # ✅ This ensures the table is created

        # Insert data into the table
        for row in sheet.iter_rows(min_row=2, values_only=True):  # Skip header row
            row_data = {}
            for i in range(len(sanitized_headers)):
                value = row[i]
                # Konverter tom tekst eller tom celle til None
                if value in ("", None):
                    value = None

                if isinstance(value, float) and value.is_integer():
                    value = int(value)
                row_data[sanitized_headers[i]] = value
            print(f"Row Data: {row_data}")  # Debugging: Print row data
            instance = DynamicModel(**row_data)
            db.session.add(instance)

        db.session.commit()

        # Clean up the uploaded file
        if os.path.exists(file_path):
            os.remove(file_path)

        return jsonify({'success': True, 'table_name': table_name})

    except Exception as e:
        print(f'Error processing file: {e}')
        traceback.print_exc()
        return jsonify({'error': 'Failed to process the Excel file'}), 500
